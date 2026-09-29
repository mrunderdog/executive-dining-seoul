#!/usr/bin/env python3
from __future__ import annotations

import argparse
import io
import json
import re
import sys
from collections import Counter
from pathlib import Path
from urllib.parse import urlencode

import requests
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry
from openpyxl import load_workbook
import xlrd

ROOT = Path(__file__).resolve().parents[1]
REGISTRY = ROOT / "sources" / "public_agency_registry.json"
RAW_OUT = ROOT / "data" / "raw" / "public_agency_executive_expense.json"
REPORT_JSON = ROOT / "reports" / "public-agency-executive-ingestion.json"
REPORT_MD = ROOT / "reports" / "public-agency-executive-ingestion.md"

UA = "Mozilla/5.0 (compatible; ExecutiveDining/1.0; +https://github.com/mrunderdog/executive-dining-seoul)"


def robust_session() -> requests.Session:
    s=requests.Session()
    retry=Retry(total=4,connect=4,read=3,backoff_factor=1.2,status_forcelist=(429,500,502,503,504),allowed_methods=frozenset(["GET"]))
    s.mount("https://",HTTPAdapter(max_retries=retry))
    s.mount("http://",HTTPAdapter(max_retries=retry))
    s.headers.update({"User-Agent":UA,"Referer":"https://alio.go.kr/"})
    return s

MERCHANT_ALIASES = ("사용처","업체명","상호","가맹점","집행장소","장소","결제처","거래처","대상업체")
DATE_ALIASES = ("일자","집행일","사용일","일시","집행일자","사용일자")
AMOUNT_ALIASES = ("금액","집행금액","사용금액","결제금액")
PURPOSE_ALIASES = ("집행목적","목적","집행내역","사용내역","내용","내역")
PEOPLE_ALIASES = ("인원","인원수","참석인원")
ROLE_ALIASES = ("직위","대상","사용자","집행자")


def t(v) -> str:
    return " ".join(str(v or "").split()).strip()


def norm(v) -> str:
    return re.sub(r"\s+","",t(v)).lower()


def find_col(headers: list[str], aliases: tuple[str, ...]) -> int | None:
    nh=[norm(x) for x in headers]
    for a in aliases:
        na=norm(a)
        for i,h in enumerate(nh):
            if h==na or na in h:
                return i
    return None


def parse_amount(v) -> int:
    if isinstance(v,(int,float)):
        return int(v)
    s=re.sub(r"[^0-9.-]","",t(v))
    if not s:
        return 0
    try:
        return int(float(s))
    except ValueError:
        return 0


def parse_date(v) -> str:
    if hasattr(v,"strftime"):
        try:
            return v.strftime("%Y-%m-%d")
        except Exception:
            pass
    s=t(v)
    m=re.search(r"(20\d{2})[./-](\d{1,2})[./-](\d{1,2})",s)
    if m:
        return f"{int(m.group(1)):04d}-{int(m.group(2)):02d}-{int(m.group(3)):02d}"
    m=re.search(r"(\d{1,2})[./-](\d{1,2})",s)
    if m:
        return f"{int(m.group(1)):02d}-{int(m.group(2)):02d}"
    return s


def report_page(session: requests.Session, apba_id: str, root_no: str) -> tuple[str,str,list[dict]]:
    url=f"https://alio.go.kr/item/itemReportTerm.do?apbaId={apba_id}&disclosureNo=&reportFormRootNo={root_no}"
    html=session.get(url,timeout=45).text
    m=re.search(r'disclosureNo\s*:\s*"([^"]+)"',html)
    if not m:
        raise RuntimeError("disclosureNo_not_found")
    disclosure=m.group(1)
    files=[]
    for no,name in re.findall(r'<option value="(\d+)">\s*([^<]+?)\s*</option>',html,re.I|re.S):
        if "선택" in name:
            continue
        files.append({"file_no":no,"name":t(name)})
    return url,disclosure,files


def workbook_rows(content: bytes) -> list[tuple[str, list[list]]]:
    out=[]
    if content.startswith(b"PK"):
        wb=load_workbook(io.BytesIO(content),data_only=True,read_only=True)
        for ws in wb.worksheets:
            out.append((ws.title, [[x for x in r] for r in ws.iter_rows(values_only=True)]))
        return out
    if content.startswith(b"\xd0\xcf\x11\xe0"):
        wb=xlrd.open_workbook(file_contents=content)
        for ws in wb.sheets():
            rows=[]
            for ridx in range(ws.nrows):
                vals=[]
                for cidx in range(ws.ncols):
                    cell=ws.cell(ridx,cidx)
                    v=cell.value
                    if cell.ctype==xlrd.XL_CELL_DATE:
                        try:
                            v=xlrd.xldate_as_datetime(v,wb.datemode)
                        except Exception:
                            pass
                    vals.append(v)
                rows.append(vals)
            out.append((ws.name,rows))
        return out
    raise ValueError("unsupported_workbook_format")


def read_workbook(content: bytes) -> list[dict]:
    parsed=[]
    for sheet_name,rows in workbook_rows(content):
        if not rows:
            continue
        header_idx=None
        mapping=None
        for i,row in enumerate(rows[:35]):
            headers=[t(x) for x in row]
            merchant=find_col(headers,MERCHANT_ALIASES)
            if merchant is None:
                continue
            mapping={
                "merchant":merchant,
                "date":find_col(headers,DATE_ALIASES),
                "amount":find_col(headers,AMOUNT_ALIASES),
                "purpose":find_col(headers,PURPOSE_ALIASES),
                "people":find_col(headers,PEOPLE_ALIASES),
                "role":find_col(headers,ROLE_ALIASES),
            }
            header_idx=i
            break
        if header_idx is None or not mapping:
            continue
        for row in rows[header_idx+1:]:
            if mapping["merchant"]>=len(row):
                continue
            merchant=t(row[mapping["merchant"]])
            if not merchant or merchant in {"합계","계","소계","총계","-"}:
                continue
            def val(k):
                idx=mapping[k]
                return row[idx] if idx is not None and idx<len(row) else None
            parsed.append({
                "sheet":sheet_name,
                "merchant":merchant,
                "date":parse_date(val("date")),
                "amount":parse_amount(val("amount")),
                "purpose":t(val("purpose")),
                "people":parse_amount(val("people")),
                "role":t(val("role")),
            })
    return parsed


def inspect_workbook(content: bytes) -> dict:
    sheets=workbook_rows(content)
    terms=Counter()
    nonempty=0
    header_hits=[]
    for sheet_name,rows in sheets:
        for ridx,row in enumerate(rows,1):
            vals=[t(x) for x in row]
            if any(vals):
                nonempty+=1
            for v in vals:
                nv=norm(v)
                if nv:
                    for a in MERCHANT_ALIASES:
                        if norm(a) in nv:
                            terms[a]+=1
                            header_hits.append({"sheet":sheet_name,"row":ridx,"value":v})
    return {"sheets":[x[0] for x in sheets],"nonempty_rows":nonempty,"merchant_header_terms":dict(terms),"merchant_header_hits":header_hits[:10]}

def meal_like(row: dict) -> bool:
    s=" ".join([t(row.get("merchant")),t(row.get("purpose"))])
    if any(x in s for x in ("경조","화환","상품권","기념품","주유","주차","택시","교통","숙박")):
        return False
    return any(x in s for x in ("식사","오찬","만찬","간담","회의","업무협의","협의","식당","한식","일식","중식","횟집","갈비","카페","커피")) or bool(t(row.get("merchant")))


def main():
    ap=argparse.ArgumentParser()
    ap.add_argument("--latest-files",type=int,default=2)
    ap.add_argument("--agency")
    args=ap.parse_args()

    registry=json.loads(REGISTRY.read_text(encoding="utf-8"))
    root_no=str(registry.get("report_form_root_no") or "20701")
    agencies=[a for a in registry.get("agencies",[]) if a.get("enabled")]
    if args.agency:
        agencies=[a for a in agencies if a.get("id")==args.agency]
    session=robust_session()

    status=[]
    raw_rows=[]
    for agency in agencies:
        item={"id":agency["id"],"name":agency["name"],"apba_id":agency["apba_id"],"state":"ERROR","files":[]}
        try:
            page_url,disclosure,files=report_page(session,agency["apba_id"],root_no)
            item["report_url"]=page_url
            item["disclosure_no"]=disclosure
            item["available_files"]=len(files)
            def file_year(x):
                years=re.findall(r"20\d{2}",x.get("name") or "")
                return max([int(y) for y in years],default=0)
            selected=sorted(files,key=lambda x:(file_year(x),int(x.get("file_no") or 0)),reverse=True)[:max(1,args.latest_files)]
            merchant_total=0
            for f in selected:
                url="https://alio.go.kr/download/file.json?"+urlencode({"f":f["file_no"],"d":disclosure})
                resp=session.get(url,timeout=60)
                file_info={"name":f["name"],"file_no":f["file_no"],"status":resp.status_code,"bytes":len(resp.content)}
                if resp.status_code!=200 or not (resp.content.startswith(b"PK") or resp.content.startswith(b"\xd0\xcf\x11\xe0")):
                    file_info["state"]="DOWNLOAD_FAILED"
                    item["files"].append(file_info)
                    continue
                inspection=inspect_workbook(resp.content)
                parsed=read_workbook(resp.content)
                file_info["inspection"]=inspection
                file_info["merchant_rows"]=len(parsed)
                file_info["state"]="MERCHANT_LEVEL" if parsed else "AGGREGATE_ONLY"
                merchant_total+=len(parsed)
                for row in parsed:
                    if not meal_like(row):
                        continue
                    row.update({
                        "agency_id":agency["id"],
                        "agency_name":agency["name"],
                        "source_file":f["name"],
                        "source_url":page_url,
                        "disclosure_no":disclosure,
                    })
                    raw_rows.append(row)
                item["files"].append(file_info)
            item["merchant_rows"]=merchant_total
            item["state"]="MERCHANT_LEVEL" if merchant_total else "TRACK_ONLY_AGGREGATE"
        except Exception as exc:
            item["error"]=f"{type(exc).__name__}: {exc}"
        status.append(item)
        print(json.dumps(item,ensure_ascii=False))

    raw_doc={
        "schema_version":1,
        "source":"ALIO 기관장 업무추진비",
        "rows":raw_rows,
        "agencies":status,
    }
    RAW_OUT.parent.mkdir(parents=True,exist_ok=True)
    RAW_OUT.write_text(json.dumps(raw_doc,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")

    report={
        "schema_version":1,
        "agency_count":len(status),
        "merchant_level_agencies":sum(x.get("state")=="MERCHANT_LEVEL" for x in status),
        "aggregate_only_agencies":sum(x.get("state")=="TRACK_ONLY_AGGREGATE" for x in status),
        "errors":sum(x.get("state")=="ERROR" for x in status),
        "merchant_rows":len(raw_rows),
        "agencies":status,
    }
    REPORT_JSON.parent.mkdir(parents=True,exist_ok=True)
    REPORT_JSON.write_text(json.dumps(report,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    md=[
        "# 공공기관 기관장 업무추진비 수집 현황",
        "",
        f"- 대상 기관: {report['agency_count']}개",
        f"- 식당/사용처 수준 확인: {report['merchant_level_agencies']}개",
        f"- 집계형만 확인: {report['aggregate_only_agencies']}개",
        f"- 오류: {report['errors']}개",
        f"- 지도 후보 원자료 행: {report['merchant_rows']}건",
        "",
        "| 기관 | 상태 | merchant rows | 비고 |",
        "|---|---:|---:|---|",
    ]
    for x in status:
        note=x.get("error") or ", ".join(f"{f.get('name')}:{f.get('state')}" for f in x.get("files",[]))
        md.append(f"| {x['name']} | {x['state']} | {x.get('merchant_rows',0)} | {note} |")
    REPORT_MD.write_text("\n".join(md)+"\n",encoding="utf-8")

    print(json.dumps(report,ensure_ascii=False))


if __name__=="__main__":
    main()
