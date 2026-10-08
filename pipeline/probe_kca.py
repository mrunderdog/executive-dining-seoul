#!/usr/bin/env python3
from __future__ import annotations
import io,ssl,urllib.request
import openpyxl

URL="https://www.kca.go.kr/kca/board/download.do?menukey=5152&fno=10055926&bid=00000030&did=1004560361"
UA="Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 Chrome/153 Safari/537.36"
req=urllib.request.Request(URL,headers={"User-Agent":UA,"Referer":"https://www.kca.go.kr/kca/sub.do?menukey=5152&mode=view&no=1004560361&page=1"})
with urllib.request.urlopen(req,timeout=20,context=ssl._create_unverified_context()) as r:
    blob=r.read()
    print("DOWNLOAD",len(blob),r.geturl(),r.headers.get("content-type"),r.headers.get("content-disposition"),blob[:8])
wb=openpyxl.load_workbook(io.BytesIO(blob),data_only=True,read_only=True)
print("SHEETS",wb.sheetnames)
for ws in wb.worksheets:
    print("\nSHEET",ws.title)
    for ri,row in enumerate(ws.iter_rows(min_row=1,max_row=min(ws.max_row,30),values_only=True),1):
        vals=["" if v is None else str(v) for v in row]
        if any(v.strip() for v in vals):
            print(ri,vals)
