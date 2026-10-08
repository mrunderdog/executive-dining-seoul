#!/usr/bin/env python3
from __future__ import annotations
import io,json,urllib.parse,urllib.request
from openpyxl import load_workbook
BASE="https://www.fowi.or.kr"
REF=BASE+"/user/publication/coreView.do?coreId=22&menu=core_"
UA="Mozilla/5.0"
def get(url):
    req=urllib.request.Request(url,headers={"User-Agent":UA,"Referer":REF,"X-Requested-With":"XMLHttpRequest"})
    with urllib.request.urlopen(req,timeout=20) as r:return r.read(),r.geturl(),r.headers
raw,_,_=get(BASE+"/selectGenerateDownloadLink.do?fileId=30857")
link=json.loads(raw.decode("utf-8"))["result"]
blob,u,h=get(urllib.parse.urljoin(BASE,link))
print("DOWNLOAD",len(blob),u,h.get("content-type"),h.get("content-disposition"),blob[:8])
wb=load_workbook(io.BytesIO(blob),read_only=True,data_only=True)
print("SHEETS",wb.sheetnames)
for ws in wb.worksheets:
    print("\nSHEET",ws.title,ws.max_row,ws.max_column)
    for ri,row in enumerate(ws.iter_rows(values_only=True),1):
        vals=[" ".join(str(v or "").split()) for v in row]
        if any(vals):
            print(ri,vals)
        if ri>=50:break
