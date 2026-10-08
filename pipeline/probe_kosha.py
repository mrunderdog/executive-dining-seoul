#!/usr/bin/env python3
from __future__ import annotations
import io,json,urllib.parse,urllib.request
from openpyxl import load_workbook

BASE="https://www.kosha.or.kr"
PROCESS=BASE+"/api/compn24/auth/stdtboard/process.do"
FILE=BASE+"/api/compn24/auth/stdtboard/fileDownload.do"
UA="Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 Chrome/153 Safari/537.36"

def call(bbs,service,data):
    obj={
      "common":{"frontInfo":{"viewId":"","menuId":"","siteId":""},"frontAuthKey":"","auth":{},"securityInfo":{},
        "data":{"pagingInfo":None,"whereId":None,"tboard":{"systemCd":"50","channel":"web","bbsId":bbs,"bbsGrpId":"","serviceId":service}}},
      "service":{"info":{"id":"","type":""},"data":data}
    }
    raw=json.dumps(obj,ensure_ascii=False,separators=(",",":"))
    double=urllib.parse.quote(urllib.parse.quote(raw,safe=""),safe="")
    body=("_JSON="+double).encode()
    req=urllib.request.Request(PROCESS,data=body,headers={
      "User-Agent":UA,"Content-Type":"application/x-www-form-urlencoded; charset=UTF-8",
      "Origin":BASE,"Referer":BASE+"/esg/management-disclosure/self-disclosure/executive-expenses/institution-leader"
    })
    with urllib.request.urlopen(req,timeout=25) as r:
        b=r.read();print("CALL",service,r.status,r.headers.get("content-type"),len(b))
    return json.loads(b.decode("utf-8"))

bbs="B2025021400028"
pst="20260916172624OU66X6"
res=call(bbs,"fileDownloadAll",{"pstNo":pst,"artclNo":"D080100001","bbsAtcflNo":"all","pstNm":"","fileNm":""})
print("RESP_KEYS",res.keys())
fd=res["response"]["fileDownList"][0]
print("FILEINFO",json.dumps({k:(v[:80]+"..." if isinstance(v,str) and len(v)>100 else v) for k,v in fd.items()},ensure_ascii=False))
url=FILE+"?"+urllib.parse.urlencode({"data":fd["data"],"key":fd["key"]})
req=urllib.request.Request(url,headers={"User-Agent":UA,"Referer":BASE+"/"})
with urllib.request.urlopen(req,timeout=25) as r:
    blob=r.read()
    print("DOWNLOAD",r.status,len(blob),r.headers.get("content-type"),r.headers.get("content-disposition"),blob[:8])
wb=load_workbook(io.BytesIO(blob),data_only=True)
for ws in wb.worksheets:
    print("SHEET",ws.title,ws.max_row,ws.max_column)
    for row in ws.iter_rows(min_row=1,max_row=min(ws.max_row,24),values_only=True):
        print(repr(row))
