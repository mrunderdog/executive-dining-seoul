"""Audit official expense attachments for address columns and exact source rows.

No address is written to public data. All matches remain human-review candidates.
"""
from __future__ import annotations
import argparse
from collections import defaultdict, Counter
import hashlib
from io import BytesIO
from html.parser import HTMLParser
import json
from pathlib import Path
import re
from urllib.parse import urlparse, urljoin
from urllib.request import Request, build_opener, HTTPRedirectHandler
import zipfile
import xml.etree.ElementTree as ET

ROOT = Path(__file__).resolve().parents[1]
FILES = {
    "공기업·공공기관": "public-enterprise-candidates.json",
    "중앙정부": "central-executive-candidates.json",
    "법조·법률행정": "justice-leadership-candidates.json",
}
SAFE_HOSTS = {
    "www.kamco.or.kr", "www.motie.go.kr", "www.moel.go.kr",
    "www.moj.go.kr", "www.moleg.go.kr", "www.unikorea.go.kr",
    "www.ksure.or.kr", "www.koreapost.go.kr", "www.hf.go.kr",
}
NAME_FIELDS = {"상호", "상호명", "업체명", "가맹점명", "거래처", "거래처명", "사용처", "업소명", "집행장소", "사용장소"}
ADDRESS_FIELDS = {"주소", "소재지", "도로명주소", "사업장주소", "업체주소", "가맹점주소", "거래처주소", "사용처주소", "집행장소주소", "사용장소주소", "영업장주소"}
MAX_BYTES = 8_000_000

def norm(value):
    return re.sub(r"[\s·()（）]", "", str(value or "")).lower()

class RestrictedRedirect(HTTPRedirectHandler):
    def redirect_request(self, request, fp, code, msg, headers, newurl):
        p = urlparse(newurl)
        if p.scheme != "https" or p.hostname not in SAFE_HOSTS:
            raise ValueError("Unsafe redirected host")
        return super().redirect_request(request, fp, code, msg, headers, newurl)

def fetch(url):
    p = urlparse(url)
    if p.scheme != "https" or p.hostname not in SAFE_HOSTS or p.username or p.password:
        raise ValueError("URL not on official allowlist")
    opener = build_opener(RestrictedRedirect)
    with opener.open(Request(url, headers={"User-Agent":"OfficialVenueSourceAudit/1.0"}), timeout=13) as resp:
        content = resp.read(MAX_BYTES + 1)
        content_type = resp.headers.get("Content-Type","")
    if len(content)>MAX_BYTES:
        raise ValueError("Over 8 MB download guard")
    return content, content_type


class OfficialAttachmentLinks(HTMLParser):
    """Collect review-only links labeled with a document filename.

    Posting pages are untrusted HTML. Parsing only produces candidate strings:
    it must never make an extra request or promote evidence.
    """
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.href = None
        self.label = []
        self.links = []

    def handle_starttag(self, tag, attrs):
        if tag.lower() == "a":
            self.href = dict(attrs).get("href") or ""
            self.label = []

    def handle_data(self, data):
        if self.href is not None:
            self.label.append(data)

    def handle_endtag(self, tag):
        if tag.lower() == "a" and self.href is not None:
            label = "".join(self.label)
            if re.search(r"\.(?:xlsx?|csv|hwpx?|pdf)(?:\s|$|\b)", label, re.I):
                self.links.append((self.href, label[:160]))
            self.href = None
            self.label = []


def attachment_candidates(html, posting_url):
    """Return at most four same-host HTTPS review links, no network requests.

    An accepted link is not evidence that an attachment exists or that its
    contents match an expense merchant; manual validation is still required.
    """
    origin = urlparse(posting_url)
    if origin.scheme != "https" or origin.hostname not in SAFE_HOSTS or origin.username or origin.password:
        return []
    parser = OfficialAttachmentLinks()
    parser.feed(html)
    result = []
    seen = set()
    for href, label in parser.links[:120]:
        absolute = urljoin(posting_url, href)
        parsed = urlparse(absolute)
        if (parsed.scheme != "https" or parsed.hostname != origin.hostname
                or parsed.username or parsed.password or not parsed.path
                or absolute in seen):
            continue
        seen.add(absolute)
        result.append(absolute)
        if len(result) == 4:
            break
    return result

def source_index():
    by = defaultdict(lambda: defaultdict(set))
    for cohort, filename in FILES.items():
        doc = json.loads((ROOT/"reports"/filename).read_text(encoding="utf-8"))
        for candidate in doc.get("candidates",[]):
            if str(candidate.get("address") or "").strip():
                continue
            for event in candidate.get("recent",[]):
                url = str(event.get("source") or "").strip()
                if url.startswith("https://"):
                    by[cohort][url].add(str(candidate.get("merchant") or "").strip())
    return by

def inspect_xlsx(raw, merchant_names):
    from openpyxl import load_workbook
    wb = load_workbook(BytesIO(raw), read_only=True, data_only=True)
    result = {"sheets":[], "matched_rows":[]}
    norm_names = {norm(n):n for n in merchant_names if n}
    try:
        for ws in wb.worksheets[:15]:
            rows = ws.iter_rows(values_only=True)
            early=[]
            for _ in range(22):
                row = next(rows, None)
                if row is None: break
                early.append(row)
            headers=[]
            for idx, row in enumerate(early):
                values=[norm(v) for v in row]
                nc=[i for i,v in enumerate(values) if v in {norm(x) for x in NAME_FIELDS}]
                ac=[i for i,v in enumerate(values) if v in {norm(x) for x in ADDRESS_FIELDS}]
                if nc and ac: headers.append((idx,nc[0],ac[0]))
            sheet={"sheet":ws.title,"address_field_present":bool(headers),"examined_header_rows":len(early),"exact_candidate_rows":0}
            if headers:
                hdr,ncol,acol=headers[0]
                for excel_row, row in enumerate(early[hdr+1:]+list(rows),start=hdr+2):
                    if max(ncol,acol)>=len(row):continue
                    merchant=norm_names.get(norm(row[ncol]))
                    address=str(row[acol] or "").strip()
                    if not merchant or not address:continue
                    # Never accept locality-only or organizational addresses.
                    if not re.search(r"(?:로|길)\s*\d+|(?:동|읍|면)\s*\d+",address):
                        continue
                    result["matched_rows"].append({
                        "merchant":merchant,"candidate_address":address[:180],
                        "sheet":ws.title,"excel_row":excel_row,
                        "review_status":"OFFICIAL_SOURCE_ROW_CANDIDATE_NOT_APPROVED",
                        "payment_branch_verified":False,"personal_attendance_verified":False
                    })
                    sheet["exact_candidate_rows"]+=1
            result["sheets"].append(sheet)
    finally:
        wb.close()
    return result

def inspect_hwpx(raw):
    # Presence of text is NOT a relational (merchant,address) match.
    with zipfile.ZipFile(BytesIO(raw)) as z:
        names=[n for n in z.namelist() if re.match(r"^Contents/section\d+\.xml$",n)]
        if not names:return {"section_count":0,"address_label_present":False}
        count=0
        for n in names[:30]:
            if z.getinfo(n).file_size>3_000_000: continue
            tree=ET.fromstring(z.read(n))
            text=" ".join(tree.itertext())
            if any(label in text for label in ("사업장 주소","사업장주소","가맹점 주소","업체주소","도로명주소")):
                count+=1
        return {"section_count":len(names),"sections_with_address_label":count,
            "address_label_present":count>0,"row_level_address_verified":False}

def inspect_doc(url, merchants):
    out={"source_url":url,"affected_merchant_count":len(merchants),
         "status":"UNPROCESSED","attachment_sha256":None,
         "exact_official_source_row_candidates":[]}
    try:
        raw, content_type = fetch(url)
        out["attachment_sha256"]=hashlib.sha256(raw).hexdigest()
        out["byte_length"]=len(raw)
        out["response_content_type"]=content_type[:100]
        if zipfile.is_zipfile(BytesIO(raw)):
            with zipfile.ZipFile(BytesIO(raw)) as z:
                members=set(z.namelist())
            if "xl/workbook.xml" in members:
                audit=inspect_xlsx(raw,merchants)
                out["sheets"]=audit["sheets"]
                out["exact_official_source_row_candidates"]=audit["matched_rows"]
                out["status"]="XLSX_CANDIDATES_REVIEW" if audit["matched_rows"] else "XLSX_NO_EXACT_ADDRESS_ROW"
            elif "mimetype" in members and "Contents/content.hpf" in members:
                out["hwpx"]=inspect_hwpx(raw)
                out["status"]="HWPX_TEXT_REVIEW_REQUIRED"
            else:
                out["status"]="ZIP_UNKNOWN_TYPE"
        elif raw[:4]==b"%PDF":
            out["status"]="PDF_REQUIRES_LAYOUT_REVIEW"
        elif raw[:8]==b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
            out["status"]="LEGACY_OFFICE_REQUIRES_PARSER"
        elif b"<html" in raw[:2000].lower() or "text/html" in content_type.lower():
            out["status"]="HTML_ATTACHMENT_DISCOVERY_REVIEW"
            out["same_host_download_candidates"]=attachment_candidates(
                raw.decode("utf-8",errors="replace"), url
            )
            out["attachment_candidates_not_downloaded"]=True
        else:
            out["status"]="UNKNOWN_FILETYPE"
    except Exception as exc:
        out["status"]="FETCH_OR_PARSE_FAILED"
        out["failure_type"]=type(exc).__name__
        out["failure_detail"]=str(exc)[:200]
    return out

def audit(limit):
    index=source_index()
    result={"policy":"Candidate row matches require manual review; do not write addresses to public entity or assert card merchant/person attendance.",
            "max_per_cohort":limit,"cohorts":[]}
    for cohort, urls in index.items():
        # Highest distinct affected-merchant leverage, then stable URL.
        chosen=sorted(urls,key=lambda u:(-len(urls[u]),u))[:limit]
        audits=[inspect_doc(u,urls[u]) for u in chosen]
        statuses=Counter(a["status"] for a in audits)
        result["cohorts"].append({"cohort":cohort,"total_source_urls":len(urls),
          "audited_source_urls":len(audits),"result_status_counts":dict(statuses),"documents":audits})
    return result

if __name__=="__main__":
    cli=argparse.ArgumentParser()
    cli.add_argument("--limit",type=int,default=3)
    cli.add_argument("--output",default="reports/official-source-address-field-audit.json")
    opts=cli.parse_args()
    if not 1<=opts.limit<=8:cli.error("--limit must be in 1..8")
    doc=audit(opts.limit)
    path=ROOT/opts.output
    path.parent.mkdir(parents=True,exist_ok=True)
    path.write_text(json.dumps(doc,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    print(json.dumps([{"cohort":c["cohort"],"counts":c["result_status_counts"],
        "exact_candidates":sum(len(d["exact_official_source_row_candidates"]) for d in c["documents"])}
        for c in doc["cohorts"]],ensure_ascii=False))
