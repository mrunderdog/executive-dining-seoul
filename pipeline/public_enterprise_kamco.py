#!/usr/bin/env python3
from __future__ import annotations

import html
import io
import json
import re
import time
import urllib.parse
import urllib.request
from collections import deque
from datetime import date, datetime
from pathlib import Path
from urllib.parse import parse_qs, urlparse

import openpyxl

from public_enterprise_discovery import decode, parse_links

ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / "reports" / "public-enterprise-discovery.json"
KEY = "kamco"
INSTITUTION = "한국자산관리공사"
SEED = "https://www.kamco.or.kr/portal/bbs/view.do?bIdx=22548&mId=0601060603&ptIdx=479"
LISTING = "https://www.kamco.or.kr/portal/bbs/list.do?mId=0601060603&ptIdx=479"
UA = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/153.0.0.0 Safari/537.36"
ATTACH_RE = re.compile(
    r"onclick=[\"']fn_egov_downFile\(\s*[\"']([^\"']+)[\"']\s*,\s*[\"']?(\d+)[\"']?\s*\).*?<span[^>]*>(.*?)</span>",
    re.I | re.S,
)


def _bidx(url: str) -> str:
    try:
        return (parse_qs(urlparse(url).query).get("bIdx") or [""])[0]
    except Exception:
        return ""


def _plain(value) -> str:
    return " ".join(str(value or "").replace("\xa0", " ").split()).strip()


def _fetch_page(url: str, attempts: int = 1) -> str:
    last = None
    for attempt in range(attempts):
        try:
            req = urllib.request.Request(url, headers={
                "User-Agent": UA,
                "Accept": "text/html,application/xhtml+xml,*/*;q=0.8",
                "Accept-Language": "ko-KR,ko;q=0.9,en;q=0.7",
            })
            with urllib.request.urlopen(req, timeout=7) as r:
                return decode(r.read(), r.headers.get_content_charset())
        except Exception as e:
            last = e
            if attempt + 1 < attempts:
                time.sleep(0.5 * (attempt + 1))
    raise last


def _date(value) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    s = _plain(value)
    m = re.search(r"(20\d{2})[-./년\s]+(\d{1,2})[-./월\s]+(\d{1,2})", s)
    if not m:
        return ""
    try:
        return date(int(m.group(1)), int(m.group(2)), int(m.group(3))).isoformat()
    except ValueError:
        return ""


def _int(value):
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return int(round(value))
    s = re.sub(r"[^0-9.-]", "", str(value))
    try:
        return int(round(float(s))) if s else None
    except Exception:
        return None


def _role(executor: str, sheet: str) -> str:
    if sheet == "기관장":
        return "기관장"
    s = re.sub(r"\s+", "", _plain(executor))
    if "부사장" in s:
        return "부사장"
    if "상임감사" in s or "감사" in s:
        return "감사"
    if "이사" in s:
        return "이사"
    return "임원"


def _download_url(file_id: str, file_sn: str) -> str:
    return "https://www.kamco.or.kr/cmm/fms/FileDown.do?" + urllib.parse.urlencode({
        "atchFileId": file_id,
        "fileSn": file_sn,
    })


def _download(url: str, parent: str, attempts: int = 1) -> bytes:
    last = None
    for attempt in range(attempts):
        try:
            req = urllib.request.Request(url, headers={
                "User-Agent": UA,
                "Referer": parent,
                "Accept": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,application/vnd.ms-excel,*/*",
                "Accept-Language": "ko-KR,ko;q=0.9,en;q=0.7",
            })
            with urllib.request.urlopen(req, timeout=12) as r:
                blob = r.read()
            if blob[:2] != b"PK":
                raise ValueError(f"attachment is not xlsx ({len(blob)} bytes)")
            return blob
        except Exception as e:
            last = e
            if attempt + 1 < attempts:
                time.sleep(0.75 * (attempt + 1))
    raise last


def _sheet_rows(blob: bytes, attachment: dict) -> list[dict]:
    wb = openpyxl.load_workbook(io.BytesIO(blob), data_only=True, read_only=True)
    out: list[dict] = []
    seq = attachment.get("file_id", "")
    source_url = attachment.get("parent", "")
    for sheet_name in ("기관장", "임원"):
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        header_row = None
        headers: list[str] = []
        for ri, row in enumerate(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 12), values_only=True), start=1):
            vals = [_plain(v).replace(" ", "") for v in row]
            joined = "|".join(vals)
            if "사용일자" in joined and "사용처" in joined and "집행금액" in joined:
                header_row = ri
                headers = vals
                break
        if not header_row:
            continue

        def col(*names: str):
            for i, h in enumerate(headers):
                if any(n.replace(" ", "") in h for n in names):
                    return i
            return None

        date_i = col("사용일자")
        exec_i = col("집행자")
        purpose_i = col("집행내역")
        merchant_i = col("사용처")
        target_i = col("집행대상자")
        payment_i = col("집행구분")
        people_i = col("인원")
        amount_i = col("집행금액")
        if None in (date_i, purpose_i, merchant_i, amount_i):
            continue

        for ri, row in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
            vals = list(row)
            need = max(date_i, purpose_i, merchant_i, amount_i)
            if need >= len(vals):
                continue
            used = _date(vals[date_i])
            merchant = _plain(vals[merchant_i])
            amount = _int(vals[amount_i])
            if not used or not merchant or merchant in {"계", "합계"} or amount is None:
                continue
            executor = _plain(vals[exec_i]) if exec_i is not None and exec_i < len(vals) else ""
            purpose = _plain(vals[purpose_i])
            target = _plain(vals[target_i]) if target_i is not None and target_i < len(vals) else ""
            payment = _plain(vals[payment_i]) if payment_i is not None and payment_i < len(vals) else ""
            people = _int(vals[people_i]) if people_i is not None and people_i < len(vals) else None
            role = _role(executor, sheet_name)
            out.append({
                "source_key": KEY,
                "institution": INSTITUTION,
                "cohort": "public_enterprise_leadership",
                "role": role,
                "department": "",
                "used_date": used,
                "used_time": "",
                "merchant": merchant,
                "address": "",
                "purpose": purpose,
                "people": people,
                "amount": amount,
                "source_amount_scale": 1,
                "payment_method": payment,
                "source_category": f"{sheet_name} 업무추진비",
                "source_url": source_url,
                "source_sheet": sheet_name,
                "source_row": ri,
                "target": target,
                "row_id": f"kamco:{seq}:{sheet_name}:{ri}",
            })
    return out


def _attachments(doc: str, parent: str, years: set[int]) -> list[dict]:
    result = []
    for m in ATTACH_RE.finditer(doc):
        file_id, file_sn, raw_name = m.group(1), m.group(2), m.group(3)
        name = _plain(html.unescape(re.sub(r"<[^>]+>", " ", raw_name)))
        ym = re.search(r"(20\d{2})\D{0,8}(1[0-2]|0?[1-9])\s*월", name)
        if not ym:
            continue
        year, month = int(ym.group(1)), int(ym.group(2))
        if year not in years:
            continue
        result.append({
            "text": name,
            "url": _download_url(file_id, file_sn),
            "year": year,
            "month": month,
            "parent": parent,
            "file_id": file_id,
            "file_sn": file_sn,
        })
    return result


def _recent_detail_urls(years: set[int]) -> tuple[list[str], list[str], list[str]]:
    details: list[str] = []
    pages: list[str] = []
    errors: list[str] = []
    seen: set[str] = set()
    for page in range(1, 4):
        sep = "&" if "?" in LISTING else "?"
        url = LISTING + sep + urllib.parse.urlencode({"pageIndex": page})
        pages.append(url)
        try:
            doc = _fetch_page(url)
        except Exception as e:
            errors.append(f"listing {url}: {type(e).__name__}: {e}")
            continue
        for link in parse_links(url, doc):
            href = link.get("url") or ""
            label = _plain(link.get("text"))
            if "임원 업무추진비" not in label or not any(str(y) in label for y in years):
                continue
            if "view.do" not in href or "ptIdx=479" not in href:
                continue
            bid = _bidx(href) or href
            if bid in seen:
                continue
            seen.add(bid)
            details.append(href)
    return details, pages, errors


def discover(year: int) -> dict:
    years = {year, year - 1}
    out = {
        "key": KEY,
        "institution": INSTITUTION,
        "cohort": "public_enterprise_leadership",
        "default_role": "임원",
        "years": sorted(years),
        "pages": [],
        "attachments": [],
        "inline_rows": [],
        "inline_replace": True,
        "errors": [],
    }
    listed_details, listing_pages, listing_errors = _recent_detail_urls(years)
    out["pages"].extend(listing_pages)
    out["errors"].extend(listing_errors)
    follow_neighbors = not listed_details
    queue = deque(listed_details or [SEED])
    seen_pages: set[str] = set()
    seen_files: set[str] = set()

    while queue and len(seen_pages) < 40:
        url = queue.popleft()
        bid = _bidx(url) or url
        if bid in seen_pages:
            continue
        seen_pages.add(bid)
        out["pages"].append(url)
        try:
            doc = _fetch_page(url)
        except Exception as e:
            out["errors"].append(f"detail {url}: {type(e).__name__}: {e}")
            continue

        if follow_neighbors:
            for link in parse_links(url, doc):
                href = link.get("url") or ""
                if "view.do" in href and "ptIdx=479" in href:
                    child = _bidx(href) or href
                    if child not in seen_pages:
                        queue.append(href)

        for att in _attachments(doc, url, years):
            fid = att["file_id"]
            if fid in seen_files:
                continue
            seen_files.add(fid)
            out["attachments"].append(att)

    for att in out["attachments"]:
        try:
            out["inline_rows"].extend(_sheet_rows(_download(att["url"], att["parent"]), att))
        except Exception as e:
            out["errors"].append(f"file {att['file_id']}: {type(e).__name__}: {e}")

    out["inline_rows"].sort(key=lambda r: (r.get("used_date", ""), r.get("merchant", ""), r.get("role", "")))
    out["parseable_attachments"] = len(out["attachments"])
    out["status"] = "PARSEABLE_FOUND" if out["inline_rows"] else (
        "FILES_FOUND_UNSUPPORTED" if out["attachments"] else (
            "FETCH_FAILED" if out["errors"] and len(out["pages"]) <= 1 else "NO_FILES_FOUND"
        )
    )
    return out


def main():
    payload = json.loads(REPORT.read_text(encoding="utf-8"))
    fresh = discover(int(payload.get("year") or datetime.now().year))
    payload["sources"] = [x for x in payload.get("sources", []) if x.get("key") != KEY] + [fresh]
    REPORT.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({
        "key": KEY,
        "status": fresh["status"],
        "pages": len(fresh["pages"]),
        "attachments": len(fresh["attachments"]),
        "rows": len(fresh["inline_rows"]),
        "merchants": len({r["merchant"] for r in fresh["inline_rows"]}),
        "roles": sorted({r["role"] for r in fresh["inline_rows"]}),
        "errors": fresh["errors"][:8],
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()
