#!/usr/bin/env python3
from __future__ import annotations

import hashlib
import io
import json
import re
import urllib.request
from datetime import date, datetime, timedelta
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
RAW_DIR = ROOT / "data" / "raw"
REPORTS = ROOT / "reports"
UA = "ExecutiveDiningSeoul/2.3 (+https://github.com/mrunderdog/executive-dining-seoul)"

# OhmyNews documents the legacy 2024 exports as:
# 총연번, 의원번호, 의원명, 당, 당ID, 지역명, 연월일, 내역, 지출액, 사용처, 분류 항목
FILES = [
    ("21대", "2024_KAPF-21.xlsx", "https://raw.githubusercontent.com/OhmyNews/KA-money/master/2024_KAPF-21.xlsx", 64_391),
    ("22대", "2024_KAPF-22.xlsx", "https://raw.githubusercontent.com/OhmyNews/KA-money/master/2024_KAPF-22.xlsx", 70_652),
]
EXPECTED_MIN_TOTAL = 130_000

COL = {
    "serial": 0,
    "member_id": 1,
    "member": 2,
    "party": 3,
    "party_id": 4,
    "district": 5,
    "date": 6,
    "purpose": 7,
    "amount": 8,
    "merchant": 9,
    "category": 10,
}


def clean(v) -> str:
    return " ".join(str(v or "").replace("\n", " ").split()).strip()


def fetch(url: str) -> bytes:
    req = urllib.request.Request(url, headers={"User-Agent": UA, "Accept": "*/*"})
    with urllib.request.urlopen(req, timeout=120) as r:
        return r.read()


def parse_date(v) -> str:
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, (int, float)) and 20_000 < float(v) < 80_000:
        try:
            return (date(1899, 12, 30) + timedelta(days=int(v))).isoformat()
        except Exception:
            return ""
    s = clean(v)
    m = re.search(r"(20\d{2})[.\-/년\s]+(\d{1,2})[.\-/월\s]+(\d{1,2})", s)
    if not m:
        # Some OCR exports use YYYYMMDD.
        m = re.search(r"\b(20\d{2})(\d{2})(\d{2})\b", s)
    if m:
        try:
            return date(int(m.group(1)), int(m.group(2)), int(m.group(3))).isoformat()
        except ValueError:
            return ""
    return ""


def parse_amount(v):
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return int(round(v))
    s = re.sub(r"[^0-9.-]", "", clean(v))
    try:
        return int(round(float(s))) if s else None
    except ValueError:
        return None


def val(row, key):
    i = COL[key]
    return row[i] if i < len(row) else None


def parse_sheet(ws, assembly: str, source_name: str, source_url: str):
    out = []
    carry = {"member_id": "", "member": "", "party": "", "party_id": "", "district": ""}
    nonempty = 0
    valid_2024 = 0
    for ri, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if not any(x not in (None, "") for x in row):
            continue
        nonempty += 1

        # Public spreadsheets sometimes visually merge/repeat 의원 metadata.
        for k in carry:
            x = clean(val(row, k))
            if x and x not in {"의원번호", "의원명", "당", "당ID", "지역명"}:
                carry[k] = x

        used_date = parse_date(val(row, "date"))
        amt = parse_amount(val(row, "amount"))
        merchant = clean(val(row, "merchant"))
        purpose = clean(val(row, "purpose"))
        category = clean(val(row, "category"))

        # This dataset is a 2024 export. Requiring a 2024 date avoids titles/totals
        # and lets us safely use positional parsing instead of fragile header inference.
        if not used_date.startswith("2024-"):
            continue
        valid_2024 += 1
        if amt is None or not merchant or not carry["member"]:
            continue
        if merchant in {"사용처", "합계", "총계"}:
            continue

        r = {
            "assembly": assembly,
            "member_id": carry["member_id"],
            "member": carry["member"],
            "party": carry["party"],
            "party_id": carry["party_id"],
            "district": carry["district"],
            "used_date": used_date,
            "purpose": purpose,
            "amount": amt,
            "merchant": merchant,
            "category": category,
            "cohort": "national_legislator",
            "source_url": source_url,
            "source_file": source_name,
            "source_sheet": ws.title,
            "source_row": ri,
        }
        r["row_id"] = hashlib.sha256(
            "|".join(str(r.get(k, "")) for k in ("assembly", "member_id", "member", "used_date", "merchant", "amount", "source_file", "source_sheet", "source_row")).encode("utf-8")
        ).hexdigest()[:20]
        out.append(r)

    return out, {"sheet": ws.title, "max_row": ws.max_row, "max_column": ws.max_column, "nonempty_rows": nonempty, "date_2024_rows": valid_2024, "parsed_rows": len(out)}


def main():
    import openpyxl

    RAW_DIR.mkdir(exist_ok=True)
    REPORTS.mkdir(exist_ok=True)
    all_rows = []
    files = []
    errors = []

    for assembly, name, url, documented_count in FILES:
        try:
            blob = fetch(url)
            wb = openpyxl.load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
            workbook_rows = []
            sheets = []
            for ws in wb.worksheets:
                rows, diag = parse_sheet(ws, assembly, name, url)
                workbook_rows.extend(rows)
                sheets.append(diag)
            files.append({
                "assembly": assembly,
                "name": name,
                "url": url,
                "bytes": len(blob),
                "documented_count": documented_count,
                "parsed_rows": len(workbook_rows),
                "sheets": sheets,
            })
            all_rows.extend(workbook_rows)
        except Exception as e:
            errors.append({"assembly": assembly, "name": name, "url": url, "error": f"{type(e).__name__}: {e}"})

    unique = {r["row_id"]: r for r in all_rows}
    rows = sorted(unique.values(), key=lambda r: (r.get("used_date") or "", r.get("member") or "", r["row_id"]))

    payload = {
        "schema_version": 3,
        "source": "OhmyNews/KA-money",
        "upstream": "Central Election Commission political-fund accounting reports obtained via information disclosure",
        "year": 2024,
        "row_count": len(rows),
        "rows": rows,
        "files": files,
        "errors": errors,
    }
    (RAW_DIR / "national_legislator_2024_expense.json").write_text(
        json.dumps(payload, ensure_ascii=False, separators=(",", ":")), encoding="utf-8"
    )

    md = [
        "# National legislator political-fund ingestion — 2024",
        "",
        "> Source lineage: Central Election Commission accounting reports → information-disclosure PDFs → OhmyNews/Kyunghyang/Newstapa OCR/normalization → this project.",
        "",
        f"- Rows: **{len(rows):,}**",
        f"- Workbooks: **{len(files)}**",
        f"- Errors: **{len(errors)}**",
        "",
        "## Workbook diagnostics",
        "",
    ]
    for f in files:
        md.append(f"- {f['assembly']} `{f['name']}`: parsed **{f['parsed_rows']:,}** / documented **{f['documented_count']:,}**")
        for s in f.get("sheets", []):
            md.append(
                f"  - `{s['sheet']}` rows={s['max_row']:,}, cols={s['max_column']}, "
                f"2024-date={s['date_2024_rows']:,}, parsed={s['parsed_rows']:,}"
            )
    if errors:
        md += ["", "## Errors", ""]
        for e in errors:
            md.append(f"- {e}")
    (REPORTS / "national-legislator-ingestion.md").write_text("\n".join(md) + "\n", encoding="utf-8")

    print(json.dumps({"rows": len(rows), "files": len(files), "errors": len(errors)}, ensure_ascii=False))
    if len(rows) < EXPECTED_MIN_TOTAL:
        raise SystemExit(
            f"Legislator ingestion sanity gate failed: parsed {len(rows):,} rows; "
            f"documented legacy exports total {sum(x[3] for x in FILES):,}. "
            "See reports/national-legislator-ingestion.md diagnostics."
        )


if __name__ == "__main__":
    main()
