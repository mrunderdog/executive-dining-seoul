#!/usr/bin/env python3
from __future__ import annotations
import io,urllib.request
from openpyxl import load_workbook

URL="https://main.kotsa.or.kr/common/download.do?atflIdxx=F_finninfo3799630441&atflSeqn=0"
UA="Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 Chrome/153 Safari/537.36"
req=urllib.request.Request(URL,headers={"User-Agent":UA,"Referer":"https://main.kotsa.or.kr/portal/bbs/finninfo_view.do?bbscCode=finninfo&cateCode=10&bbscSeqn=37996&pageNumb=1&menuCode=03020300"})
with urllib.request.urlopen(req,timeout=20) as r:
    blob=r.read()
    print("DOWNLOAD",len(blob),r.geturl(),r.headers.get("content-type"),r.headers.get("content-disposition"),blob[:8])
wb=load_workbook(io.BytesIO(blob),data_only=True)
for ws in wb.worksheets:
    print("SHEET",ws.title,ws.max_row,ws.max_column)
    for row in ws.iter_rows(min_row=1,max_row=min(ws.max_row,30),values_only=True):
        print(repr(row))
