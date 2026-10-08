#!/usr/bin/env python3
from __future__ import annotations
import io,urllib.request
from openpyxl import load_workbook

URL="https://www.kpx.or.kr/boardDownload.es?bid=0023&list_no=78162&seq=1"
UA="Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 Chrome/153 Safari/537.36"
req=urllib.request.Request(URL,headers={"User-Agent":UA,"Referer":"https://www.kpx.or.kr/board.es?mid=a10202010000&bid=0023&act=view&list_no=78162"})
with urllib.request.urlopen(req,timeout=20) as r:
    blob=r.read()
    print("DOWNLOAD",len(blob),r.geturl(),r.headers.get("content-type"),r.headers.get("content-disposition"),blob[:8])
wb=load_workbook(io.BytesIO(blob),read_only=True,data_only=True)
for ws in wb.worksheets:
    print("\nSHEET",ws.title)
    for i,row in enumerate(ws.iter_rows(values_only=True),1):
        vals=[" ".join(str(v or "").split()) for v in row]
        if any(vals): print(i,vals)
        if i>=50: break
