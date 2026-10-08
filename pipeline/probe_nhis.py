#!/usr/bin/env python3
from __future__ import annotations
import io,urllib.request
from openpyxl import load_workbook

URL="https://www.nhis.or.kr/announce/wbhaec11200m01.do?mode=download&articleNo=11013546&attachNo=369407"
UA="Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 Chrome/153 Safari/537.36"
req=urllib.request.Request(URL,headers={"User-Agent":UA,"Referer":"https://www.nhis.or.kr/announce/wbhaec11200m01.do"})
with urllib.request.urlopen(req,timeout=20) as r:
    blob=r.read()
    print("DOWNLOAD",len(blob),r.geturl(),r.headers.get("content-type"),r.headers.get("content-disposition"),blob[:8])
wb=load_workbook(io.BytesIO(blob),read_only=True,data_only=True)
for ws in wb.worksheets:
    print("\nSHEET",ws.title)
    for i,row in enumerate(ws.iter_rows(values_only=True),1):
        vals=[" ".join(str(v or "").split()) for v in row]
        if any(vals):
            print(i,vals)
        if i>=30:break
