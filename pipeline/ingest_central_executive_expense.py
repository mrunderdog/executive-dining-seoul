#!/usr/bin/env python3
from __future__ import annotations

import argparse
import hashlib
import io
import json
import re
import urllib.request
import zipfile
import xml.etree.ElementTree as ET
from collections import Counter
from datetime import date, datetime
from pathlib import Path

ROOT=Path(__file__).resolve().parents[1]
REPORTS=ROOT/"reports"; RAW_DIR=ROOT/"data"/"raw"
DISCOVERY=REPORTS/"central-executive-discovery.json"
UA="ExecutiveDiningSeoul/2.1 (+https://github.com/mrunderdog/executive-dining-seoul)"

ALIASES={
 "date":["사용일","사용일자","집행일","집행일자","일자","사용일시","집행일시"],
 "time":["사용시간","집행시간","시간"],
 "merchant":["집행장소","사용처","사용장소","장소","업소명","상호","가맹점명"],
 "address":["주소","소재지","집행장소주소","사용처주소"],
 "purpose":["집행목적","사용목적","목적","내용","적요"],
 "people":["인원","참석인원","대상인원","사용인원"],
 "amount":["금액","집행금액","사용금액","결제금액","지출금액"],
 "role":["직위","직책","사용자","사용자명","집행자","구분"],
 "department":["부서","부서명","소속"],
 "method":["결제방법","지급방법","결제수단"]}


def clean(v):
    if v is None:return ""
    if isinstance(v,datetime):return v.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(v,date):return v.isoformat()
    return " ".join(str(v).replace("\n"," ").split()).strip()

def nh(v): return re.sub(r"[\s()\[\]·ㆍ._/-]+","",clean(v)).lower()

def amount(v):
    if isinstance(v,(int,float)) and not isinstance(v,bool): return int(round(v))
    s=re.sub(r"[^0-9.-]","",clean(v));
    try:return int(round(float(s))) if s else None
    except:return None

def people(v):
    if isinstance(v,(int,float)) and not isinstance(v,bool): return int(v)
    m=re.search(r"(\d+)",clean(v)); return int(m.group(1)) if m else None

def pdate(v):
    if isinstance(v,datetime): return v.date().isoformat()
    if isinstance(v,date): return v.isoformat()
    s=clean(v); m=re.search(r"(20\d{2})[.\-/년\s]+(\d{1,2})[.\-/월\s]+(\d{1,2})",s)
    if m:
        try:return date(int(m.group(1)),int(m.group(2)),int(m.group(3))).isoformat()
        except:pass
    return s

def fetch(url):
    req=urllib.request.Request(url,headers={"User-Agent":UA,"Accept":"*/*"})
    with urllib.request.urlopen(req,timeout=60) as r:return r.read()

def _local(tag):
    return tag.split("}",1)[-1] if "}" in tag else tag

def hwpx_tables(blob):
    tables=[]
    with zipfile.ZipFile(io.BytesIO(blob)) as zf:
        names=sorted(n for n in zf.namelist() if n.lower().startswith("contents/section") and n.lower().endswith(".xml"))
        for sec_no,name in enumerate(names,1):
            try:
                root=ET.fromstring(zf.read(name))
            except Exception:
                continue
            tno=0
            for tbl in root.iter():
                if _local(tbl.tag)!="tbl":
                    continue
                tno+=1
                rows=[]
                for tr in tbl.iter():
                    if _local(tr.tag)!="tr":
                        continue
                    row=[]
                    for tc in tr:
                        if _local(tc.tag)!="tc":
                            continue
                        parts=[]
                        for node in tc.iter():
                            if _local(node.tag)=="t" and node.text:
                                parts.append(node.text)
                        row.append(" ".join(" ".join(parts).split()))
                    if any(clean(x) for x in row):
                        rows.append(row)
                if rows:
                    tables.append((f"section{sec_no}-table{tno}",rows))
    return tables

def rows_from(blob,name):
    lower=name.lower()
    if ".hwpx" in lower:
        tables=hwpx_tables(blob)
        if not tables:
            raise ValueError("HWPX contained no parseable tables")
        for item in tables:
            yield item
        return
    if ".xlsx" in lower:
        import openpyxl
        wb=openpyxl.load_workbook(io.BytesIO(blob),read_only=True,data_only=True)
        for ws in wb.worksheets:
            yield ws.title,[list(r) for r in ws.iter_rows(values_only=True)]
        return
    if blob.startswith(bytes.fromhex("D0CF11E0A1B11AE1")) or ".xls" in lower:
        import xlrd
        book=xlrd.open_workbook(file_contents=blob)
        for sh in book.sheets():
            yield sh.name,[sh.row_values(i) for i in range(sh.nrows)]
        return
    if ".csv" in lower:
        text=blob.decode("utf-8-sig",errors="replace")
        import csv
        yield "csv",list(csv.reader(io.StringIO(text)))
        return
    if blob.startswith(b"PK"):
        # Unknown ZIP container: try HWPX first, then XLSX.
        try:
            tables=hwpx_tables(blob)
            if tables:
                for item in tables:
                    yield item
                return
        except Exception:
            pass
        import openpyxl
        wb=openpyxl.load_workbook(io.BytesIO(blob),read_only=True,data_only=True)
        for ws in wb.worksheets:
            yield ws.title,[list(r) for r in ws.iter_rows(values_only=True)]
        return
    raise ValueError("unsupported spreadsheet/document")

def header_score(row):
    cells=[nh(x) for x in row]; matched=set()
    for field,aliases in ALIASES.items():
        if any(any(nh(a)==c or nh(a) in c for a in aliases if c) for c in cells): matched.add(field)
    score=len(matched)+(2 if "merchant" in matched else 0)+(2 if "amount" in matched else 0)+(1 if "date" in matched else 0)
    return score,matched

def find_header(rows):
    best=(-1,-1,set())
    for i,row in enumerate(rows[:60]):
        s,m=header_score(row)
        if s>best[0]:best=(s,i,m)
    return best if best[0]>=4 else None

def mapping(row):
    out={}
    for i,x in enumerate(row):
        c=nh(x)
        for field,aliases in ALIASES.items():
            if field in out:continue
            if c and any(nh(a)==c or nh(a) in c for a in aliases):out[field]=i
    return out

def cell(row,m,k):
    i=m.get(k); return row[i] if i is not None and i<len(row) else None

def normalize(rows,sheet,meta):
    h=find_header(rows)
    if not h:return [],{"sheet":sheet,"status":"NO_HEADER"}
    score,hi,_=h;m=mapping(rows[hi]);out=[]
    for ri,row in enumerate(rows[hi+1:],start=hi+2):
        if not any(clean(x) for x in row):continue
        merchant=clean(cell(row,m,"merchant")); amt=amount(cell(row,m,"amount")); d=pdate(cell(row,m,"date"))
        if not merchant and amt is None:continue
        joined=" ".join(clean(x) for x in row[:10])
        if any(k in joined for k in ("합계","총계","누계")) and not d:continue
        role=clean(cell(row,m,"role")); dept=clean(cell(row,m,"department"))
        r={"source_key":meta["key"],"institution":meta["institution"],"cohort":"central_executive","role":role,"department":dept,"used_date":d,"used_time":clean(cell(row,m,"time")),"merchant":merchant,"address":clean(cell(row,m,"address")),"purpose":clean(cell(row,m,"purpose")),"people":people(cell(row,m,"people")),"amount":amt,"payment_method":clean(cell(row,m,"method")),"source_url":meta["url"],"source_sheet":sheet,"source_row":ri}
        r["row_id"]=hashlib.sha256("|".join(str(r.get(k,"")) for k in ("source_key","role","department","used_date","merchant","amount","source_sheet","source_row")).encode()).hexdigest()[:20]
        out.append(r)
    return out,{"sheet":sheet,"status":"OK","mapping":m,"parsed_rows":len(out),"header_score":score}

def main():
    ap=argparse.ArgumentParser();ap.add_argument("--year",type=int,default=datetime.now().year);args=ap.parse_args()
    if not DISCOVERY.exists():raise SystemExit("run central_executive_discovery.py first")
    d=json.loads(DISCOVERY.read_text(encoding="utf-8")); all_rows=[]; files=[]; errors=[]
    for src in d.get("sources",[]):
        for a in src.get("attachments",[]):
            s=(a.get("text","")+" "+a.get("url","")).lower()
            if not any(ext in s for ext in (".xlsx",".xls",".csv",".hwpx")):continue
            try:
                blob=fetch(a["url"]); info={"institution":src["institution"],"key":src["key"],"url":a["url"],"bytes":len(blob),"sheets":[]}
                for sheet,rows in rows_from(blob,a.get("text") or a["url"]):
                    norm,si=normalize(rows,sheet,{"key":src["key"],"institution":src["institution"],"url":a["url"]}); all_rows.extend(norm);info["sheets"].append(si)
                files.append(info)
            except Exception as e: errors.append({"institution":src["institution"],"url":a.get("url"),"error":f"{type(e).__name__}: {e}"})
    unique={r["row_id"]:r for r in all_rows}; rows=sorted(unique.values(),key=lambda r:(r.get("used_date") or "",r.get("institution") or "",r["row_id"]))
    RAW_DIR.mkdir(exist_ok=True); out=RAW_DIR/"central_executive_expense.json"; out.write_text(json.dumps({"schema_version":1,"generated_at":datetime.now().isoformat(timespec="seconds"),"row_count":len(rows),"rows":rows,"files":files,"errors":errors},ensure_ascii=False,separators=(",",":")),encoding="utf-8")
    c=Counter(r.get("institution") or "(unknown)" for r in rows);md=["# Central executive expense ingestion","",f"- Rows: **{len(rows)}**",f"- Files: **{len(files)}**",f"- Errors: **{len(errors)}**","","## Rows by institution",""]
    for k,v in c.most_common():md.append(f"- {k}: {v}")
    (REPORTS/"central-executive-ingestion.md").write_text("\n".join(md)+"\n",encoding="utf-8")
    print(json.dumps({"rows":len(rows),"files":len(files),"errors":len(errors),"institutions":len(c)},ensure_ascii=False))

if __name__=="__main__":main()
