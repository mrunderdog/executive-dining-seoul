#!/usr/bin/env python3
from __future__ import annotations

import html, io, json, re, time, urllib.parse, urllib.request
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from pathlib import Path
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
REPORT=ROOT/"reports"/"public-enterprise-discovery.json"
KEY="nhis"
INSTITUTION="국민건강보험공단"
LISTING="https://www.nhis.or.kr/announce/wbhaec11200m01.do"
UA="Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 Chrome/153 Safari/537.36"

def fetch(url:str,binary:bool=False,referer:str=""):
    last=None
    for attempt in range(3):
        try:
            req=urllib.request.Request(url,headers={
                "User-Agent":UA,"Accept":"*/*","Accept-Language":"ko-KR,ko;q=0.9",
                "Referer":referer or LISTING,"Connection":"close",
            })
            with urllib.request.urlopen(req,timeout=20) as r:
                raw=r.read()
                return raw if binary else raw.decode(r.headers.get_content_charset() or "utf-8","replace")
        except Exception as e:
            last=e
            if attempt<2: time.sleep(1+attempt)
    raise last

def clean(v):
    return " ".join(str(v or "").replace("\n"," ").split()).strip()

def date_norm(v):
    if isinstance(v,datetime):
        return v.date().isoformat()
    s=clean(v)
    m=re.search(r"(20\d{2})[-./](\d{1,2})[-./](\d{1,2})",s)
    if not m:return ""
    try:return datetime(*map(int,m.groups())).date().isoformat()
    except ValueError:return ""

def to_int(v):
    s=re.sub(r"[^0-9.-]","",str(v or ""))
    try:return int(round(float(s)))
    except Exception:return None

def discover_files(year:int):
    years={year,year-1};items=[];pages=[];errors=[]
    for off in range(0,41,10):
        url=LISTING+("?"+urllib.parse.urlencode({"article.offset":off,"articleLimit":10,"mode":"list"}) if off else "")
        try:doc=fetch(url)
        except Exception as e:
            errors.append(f"listing {url}: {type(e).__name__}: {e}");continue
        pages.append(url)
        page_years=[]
        for row in re.findall(r"<tr\b.*?</tr>",doc,re.I|re.S):
            txt=clean(html.unescape(re.sub(r"<[^>]+>"," ",row)))
            ym=re.search(r"(20\d{2})년\s*(\d{1,2})월\s*임원\s*업무추진비",txt)
            if not ym:continue
            y,mo=int(ym.group(1)),int(ym.group(2));page_years.append(y)
            if y not in years:continue
            am=re.search(r'articleNo=(\d+)&attachNo=(\d+)',row,re.I)
            if not am:continue
            article,attach=am.groups()
            dl=LISTING+"?"+urllib.parse.urlencode({"mode":"download","articleNo":article,"attachNo":attach})
            items.append({
                "year":y,"month":mo,"text":txt,
                "url":dl,"download_url":dl,"parent":url,
                "attachment_id":f"{article}|{attach}",
            })
        if page_years and min(page_years)<min(years):break
    uniq={x["attachment_id"]:x for x in items}
    return sorted(uniq.values(),key=lambda x:(x["year"],x["month"])),pages,errors

def parse_xlsx(att):
    try:
        blob=fetch(att["download_url"],True,att.get("parent",""))
        wb=load_workbook(io.BytesIO(blob),read_only=True,data_only=True)
    except Exception as e:return [],f"xlsx {att['download_url']}: {type(e).__name__}: {e}"
    rows=[]
    for ws in wb.worksheets:
        vals_all=[list(r) for r in ws.iter_rows(values_only=True)]
        header=None
        for ri,vals in enumerate(vals_all,1):
            texts=[clean(v) for v in vals]
            if "사용일자" in texts and "사용처" in texts and "집행금액(원)" in texts:
                header={name:texts.index(name) for name in texts if name}
                continue
            if not header:continue
            def get(name):
                i=header.get(name)
                return vals[i] if i is not None and i<len(vals) else ""
            d=date_norm(get("사용일자"))
            merchant=clean(get("사용처"))
            if not d or not merchant or merchant in {"-","사용처"}:continue
            role=clean(get("집행자")) if "집행자" in header else "기관장"
            purpose=clean(get("집행내역"))
            rows.append({
                "source_key":KEY,"institution":INSTITUTION,
                "cohort":"public_enterprise_leadership",
                "role":role or ws.title,"department":"",
                "used_date":d,"used_time":"",
                "merchant":merchant,"address":"",
                "purpose":purpose,
                "target":clean(get("집행대상자")),
                "payment_method":clean(get("집행구분")),
                "people":to_int(get("인원(명)")),
                "amount":to_int(get("집행금액(원)")),
                "source_amount_scale":1,
                "source_category":ws.title,
                "source_url":att["download_url"],
                "source_sheet":ws.title,"source_row":ri,
                "row_id":f"nhis:{att['attachment_id']}:{ws.title}:{ri}",
            })
    return rows,None

def discover(year:int):
    attachments,pages,errors=discover_files(year);rows=[]
    if attachments:
        with ThreadPoolExecutor(max_workers=min(6,len(attachments))) as pool:
            fm={pool.submit(parse_xlsx,a):a for a in attachments}
            for fut in as_completed(fm):
                parsed,err=fut.result();rows.extend(parsed)
                if err:errors.append(err)
    rows.sort(key=lambda r:(r["used_date"],r["role"],r["row_id"]))
    return {
        "key":KEY,"institution":INSTITUTION,
        "cohort":"public_enterprise_leadership",
        "default_role":"이사장·감사·상임이사",
        "years":[year-1,year],"pages":pages,
        "attachments":attachments,"inline_rows":rows,"inline_replace":True,
        "errors":errors,"parseable_attachments":len(attachments),
        "status":"PARSEABLE_FOUND" if rows else ("FETCH_FAILED" if errors else "NO_FILES_FOUND"),
    }

def main():
    payload=json.loads(REPORT.read_text(encoding="utf-8"))
    fresh=discover(int(payload.get("year") or datetime.now().year))
    payload["sources"]=[x for x in payload.get("sources",[]) if x.get("key")!=KEY]+[fresh]
    REPORT.write_text(json.dumps(payload,ensure_ascii=False,indent=2),encoding="utf-8")
    print(json.dumps({
        "key":KEY,"status":fresh["status"],
        "attachments":len(fresh["attachments"]),
        "rows":len(fresh["inline_rows"]),
        "merchants":len({r["merchant"] for r in fresh["inline_rows"]}),
        "roles":sorted({r["role"] for r in fresh["inline_rows"]}),
        "date_min":min((r["used_date"] for r in fresh["inline_rows"]),default=""),
        "date_max":max((r["used_date"] for r in fresh["inline_rows"]),default=""),
        "errors":fresh["errors"][:8],
    },ensure_ascii=False))

if __name__=="__main__":main()
