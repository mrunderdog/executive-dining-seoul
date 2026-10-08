#!/usr/bin/env python3
from __future__ import annotations

import io, json, re, time, urllib.parse, urllib.request
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timedelta
from pathlib import Path
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
REPORT=ROOT/"reports"/"public-enterprise-discovery.json"
KEY="kosha"
INSTITUTION="안전보건공단"
BASE="https://www.kosha.or.kr"
PROCESS=BASE+"/api/compn24/auth/stdtboard/process.do"
FILEDOWN=BASE+"/api/compn24/auth/stdtboard/fileDownload.do"
UA="Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 Chrome/153 Safari/537.36"

SOURCES=[
    {"role":"기관장","bbs":"B2025021400028","path":"institution-leader"},
    {"role":"연구원장","bbs":"B2025021400029","path":"research-director"},
    {"role":"교육원장","bbs":"B2025021400030","path":"education-director"},
    {"role":"인증원장","bbs":"B2025021400031","path":"certification-director"},
    {"role":"스마트안전보건 기술원장","bbs":"B2025021400032","path":"tech-director"},
    {"role":"감사","bbs":"B2025021400033","path":"auditorndirector","role_code":"2250001"},
    {"role":"경영기획이사","bbs":"B2025021400033","path":"auditorndirector","role_code":"2250002"},
    {"role":"안전보건사업이사","bbs":"B2025021400033","path":"auditorndirector","role_code":"2250003"},
    {"role":"교육홍보이사","bbs":"B2025021400033","path":"auditorndirector","role_code":"2250004"},
]

def role_url(src):
    return BASE+"/esg/management-disclosure/self-disclosure/executive-expenses/"+src["path"]

def _request(req, timeout=25):
    last=None
    for attempt in range(3):
        try:
            with urllib.request.urlopen(req,timeout=timeout) as r:
                return r.read(),r.headers
        except Exception as e:
            last=e
            if attempt<2:time.sleep(1+attempt)
    raise last

def tboard_call(src, service_id, service_data):
    obj={
        "common":{
            "frontInfo":{"viewId":"","menuId":"","siteId":""},
            "frontAuthKey":"","auth":{},"securityInfo":{},
            "data":{"pagingInfo":None,"whereId":None,"tboard":{
                "systemCd":"50","channel":"web","bbsId":src["bbs"],"bbsGrpId":"",
                "serviceId":service_id
            }}
        },
        "service":{"info":{"id":"","type":""},"data":service_data}
    }
    raw=json.dumps(obj,ensure_ascii=False,separators=(",",":"))
    encoded=urllib.parse.quote(urllib.parse.quote(raw,safe=""),safe="")
    body=("_JSON="+encoded).encode()
    req=urllib.request.Request(PROCESS,data=body,headers={
        "User-Agent":UA,
        "Content-Type":"application/x-www-form-urlencoded; charset=UTF-8",
        "Origin":BASE,
        "Referer":role_url(src),
        "Accept":"application/json, text/plain, */*",
        "Connection":"close",
    })
    blob,_=_request(req)
    out=json.loads(blob.decode("utf-8"))
    if out.get("code") not in (0,"0",None):
        raise RuntimeError(f"tboard {service_id} failed: {out.get('message')}")
    return out.get("response") or {}

def list_posts(src, year):
    search=[
        {"artclNo":"E030300003","artclInptCn":str(year),"cndSeCd":"01"}
    ]
    if src.get("role_code"):
        search.append({"artclNo":"E040100003","artclInptCn":src["role_code"],"cndSeCd":"01"})
    data={
        "searchDefaultCndGrid":[{
            "curPageCo":1,"recodePageCo":100,"rowsPerPage":100,
            "pstSeCd":"1200001","atcflCntSrchYn":"Y","artclNoList":[],
            "pstNoOrder":"Y","isDesc":"Y","sortType":"01","sortOrder":"1",
            "isAddPstCn":"N"
        }],
        "searchArtclCndGrid":search
    }
    resp=tboard_call(src,"basicAccess",data)
    attached={x.get("pstNo") for x in resp.get("atcflArtclGrid",[]) if x.get("bbsAtcflYn")=="Y"}
    names={x.get("pstNo"):x.get("pstNm","") for x in resp.get("bbsPstGrid",[])}
    return [{"pst_no":p,"title":names.get(p,""),"year":year} for p in attached if p]

def file_info(src, post):
    resp=tboard_call(src,"fileDownloadAll",{
        "pstNo":post["pst_no"],"artclNo":"D080100001",
        "bbsAtcflNo":"all","pstNm":"","fileNm":""
    })
    return resp.get("fileDownList") or []

def download_file(info, referer):
    q=urllib.parse.urlencode({"data":info["data"],"key":info["key"]})
    req=urllib.request.Request(FILEDOWN+"?"+q,headers={
        "User-Agent":UA,"Referer":referer,"Accept":"*/*","Connection":"close"
    })
    blob,_=_request(req)
    return blob

def _date(v):
    if isinstance(v,datetime):return v.date().isoformat()
    if isinstance(v,(int,float)) and 30000<=float(v)<=60000:
        return (datetime(1899,12,30)+timedelta(days=float(v))).date().isoformat()
    s=" ".join(str(v or "").split())
    m=re.search(r"(20\d{2})[-./년\s]+(\d{1,2})[-./월\s]+(\d{1,2})",s)
    if not m:return ""
    try:return datetime(*map(int,m.groups())).date().isoformat()
    except ValueError:return ""

def _int(v):
    if isinstance(v,(int,float)):return int(round(float(v)))
    m=re.search(r"-?\d[\d,]*",str(v or ""))
    return int(m.group().replace(",","")) if m else None

def _clean(v):
    return " ".join(str(v or "").replace("\n"," ").split()).strip()

def _merchant(v):
    s=_clean(v)
    s=re.sub(r"\s*\(☎?\s*[^)]*\)\s*$","",s)
    return s.strip()

def parse_xlsx(blob, src, post, file_name):
    rows=[]
    wb=load_workbook(io.BytesIO(blob),data_only=True,read_only=True)
    for ws in wb.worksheets:
        header=None
        for ri,vals0 in enumerate(ws.iter_rows(values_only=True),1):
            vals=list(vals0)
            texts=[_clean(v) for v in vals]
            norms=[t.replace(" ","") for t in texts]
            if header is None:
                if ("사용일자" in norms or "집행일자" in norms) and any("사용처" in x for x in norms):
                    header=ri
                continue
            d=_date(vals[0] if vals else None)
            if not d:continue
            purpose=texts[1] if len(texts)>1 else ""
            merchant=_merchant(vals[2] if len(vals)>2 else "")
            target=texts[3] if len(texts)>3 else ""
            payment=texts[4] if len(texts)>4 else ""
            if not merchant or merchant in {"-","사용처","사용처(장소)"}:continue
            people=None
            for v in vals[5:8]:
                n=_int(v)
                if n is not None:
                    people=n
            amount=None
            for v in vals[8:]:
                n=_int(v)
                if n is not None:
                    amount=n
            rows.append({
                "source_key":KEY,"institution":INSTITUTION,
                "cohort":"public_enterprise_leadership",
                "role":src["role"],"department":"",
                "used_date":d,"used_time":"",
                "merchant":merchant,"address":"","purpose":purpose,
                "target":target,"payment_method":payment,
                "people":people,"amount":amount,"source_amount_scale":1,
                "source_category":src["role"],
                "source_url":role_url(src),
                "source_sheet":ws.title,"source_row":ri,
                "row_id":f"kosha:{src['bbs']}:{src.get('role_code','')}:{post['pst_no']}:{ws.title}:{ri}",
            })
    return rows

def fetch_post(src, post):
    out=[];atts=[];errors=[]
    try:
        infos=file_info(src,post)
        for idx,info in enumerate(infos,1):
            name=info.get("fileNm") or f"{post['pst_no']}-{idx}.xlsx"
            att_id=f"{src['bbs']}|{src.get('role_code','')}|{post['pst_no']}|{idx}"
            atts.append({
                "text":post["title"],"role":src["role"],"year":post["year"],
                "url":role_url(src),"download_url":role_url(src),
                "attachment_id":att_id,"filename":name,
            })
            blob=download_file(info,role_url(src))
            if blob.startswith(b"PK"):
                out.extend(parse_xlsx(blob,src,post,name))
            else:
                errors.append(f"unsupported file {name}: {blob[:8]!r}")
    except Exception as e:
        errors.append(f"{src['role']} {post['pst_no']}: {type(e).__name__}: {e}")
    return out,atts,errors

def discover(year):
    years=[year-1,year];posts=[];pages=[];errors=[]
    for src in SOURCES:
        pages.append(role_url(src))
        for y in years:
            try:
                for p in list_posts(src,y):
                    posts.append((src,p))
            except Exception as e:
                errors.append(f"list {src['role']} {y}: {type(e).__name__}: {e}")
    rows=[];atts=[]
    if posts:
        with ThreadPoolExecutor(max_workers=min(6,len(posts))) as pool:
            fm={pool.submit(fetch_post,src,p):(src,p) for src,p in posts}
            for fut in as_completed(fm):
                r,a,e=fut.result();rows.extend(r);atts.extend(a);errors.extend(e)
    rows.sort(key=lambda r:(r["used_date"],r["role"],r["row_id"]))
    uniq_att={x["attachment_id"]:x for x in atts}
    return {
        "key":KEY,"institution":INSTITUTION,
        "cohort":"public_enterprise_leadership",
        "default_role":"기관장·감사·상임이사·원장",
        "years":years,"pages":list(dict.fromkeys(pages)),
        "attachments":list(uniq_att.values()),"inline_rows":rows,
        "inline_replace":True,"errors":errors,
        "parseable_attachments":len(uniq_att),
        "status":"PARSEABLE_FOUND" if rows else ("FETCH_FAILED" if errors else "NO_FILES_FOUND"),
    }

def main():
    payload=json.loads(REPORT.read_text(encoding="utf-8"))
    fresh=discover(int(payload.get("year") or datetime.now().year))
    payload["sources"]=[x for x in payload.get("sources",[]) if x.get("key")!=KEY]+[fresh]
    REPORT.write_text(json.dumps(payload,ensure_ascii=False,indent=2),encoding="utf-8")
    print(json.dumps({
        "key":KEY,"status":fresh["status"],"attachments":len(fresh["attachments"]),
        "rows":len(fresh["inline_rows"]),
        "merchants":len({r["merchant"] for r in fresh["inline_rows"]}),
        "roles":sorted({r["role"] for r in fresh["inline_rows"]}),
        "date_min":min((r["used_date"] for r in fresh["inline_rows"]),default=""),
        "date_max":max((r["used_date"] for r in fresh["inline_rows"]),default=""),
        "errors":fresh["errors"][:12],
    },ensure_ascii=False))

if __name__=="__main__":main()
