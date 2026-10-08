#!/usr/bin/env python3
from __future__ import annotations

import html, io, json, re, ssl, time, urllib.parse, urllib.request
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import date, datetime
from pathlib import Path

import openpyxl

ROOT=Path(__file__).resolve().parents[1]
REPORT=ROOT/"reports"/"public-enterprise-discovery.json"
KEY="kca"
INSTITUTION="한국소비자원"
LISTING="https://www.kca.go.kr/kca/sub.do?menukey=5152&mode=list&page={page}"
UA="Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 Chrome/153 Safari/537.36"
CTX=ssl._create_unverified_context()

def fetch(url:str,binary:bool=False,referer:str=""):
    last=None
    for attempt in range(3):
        try:
            req=urllib.request.Request(url,headers={
                "User-Agent":UA,
                "Accept-Language":"ko-KR,ko;q=0.9",
                "Referer":referer or url,
                "Connection":"close",
            })
            with urllib.request.urlopen(req,timeout=20,context=CTX) as r:
                raw=r.read()
                return raw if binary else raw.decode(r.headers.get_content_charset() or "utf-8","replace")
        except Exception as e:
            last=e
            if attempt<2: time.sleep(1+attempt)
    raise last

def clean(v): return " ".join(str(v or "").replace("\n"," ").split())

def norm_date(v):
    if isinstance(v,datetime): return v.date().isoformat()
    if isinstance(v,date): return v.isoformat()
    s=clean(v)
    m=re.search(r"(20\d{2})[-./년\s]+(\d{1,2})[-./월\s]+(\d{1,2})",s)
    if not m:return ""
    try:return datetime(*map(int,m.groups())).date().isoformat()
    except ValueError:return ""

def to_int(v):
    s=re.sub(r"[^0-9.-]","",str(v or ""))
    try:return int(round(float(s)))
    except Exception:return None

def merchant_clean(v):
    s=clean(v)
    s=re.sub(r"\s*\((?:0\d{1,3}|0507)[-\d]+\)\s*$","",s)
    return s.strip()

def discover_files(year:int):
    years={year,year-1}
    files=[];pages=[];errors=[];details=[]
    urls=[LISTING.format(page=p) for p in range(1,11)]
    with ThreadPoolExecutor(max_workers=8) as pool:
        fm={pool.submit(fetch,u):u for u in urls}
        for fut in as_completed(fm):
            u=fm[fut]
            try:doc=fut.result()
            except Exception as e:
                errors.append(f"listing {u}: {type(e).__name__}: {e}");continue
            pages.append(u)
            for row in re.findall(r"<tr\b.*?</tr>",doc,re.I|re.S):
                txt=clean(html.unescape(re.sub(r"<[^>]+>"," ",row)))
                tm=re.search(r"(20\d{2})년\s*(\d{1,2})월\s*일자별\s*공개\(([^)]+)\)",txt)
                if not tm:continue
                y,m,role=int(tm.group(1)),int(tm.group(2)),tm.group(3).strip()
                if y not in years:continue
                hm=re.search(r"href=['\\"]([^'\\"]*mode=view[^'\\"]*no=(\d+)[^'\\"]*)['\\"]",row,re.I)
                if not hm:continue
                detail=urllib.parse.urljoin(u,html.unescape(hm.group(1)).replace("&amp;","&"))
                details.append((y,m,role,txt,detail,hm.group(2)))

    def inspect(item):
        y,m,role,txt,detail,did=item
        ddoc=fetch(detail)
        fm2=re.search(r"href=['\\"]([^'\\"]*board/download\.do\?[^'\\"]+)['\\"]",ddoc,re.I)
        if not fm2:return detail,None
        dl=urllib.parse.urljoin(detail,html.unescape(fm2.group(1)).replace("&amp;","&"))
        q=urllib.parse.parse_qs(urllib.parse.urlparse(dl).query)
        fid=(q.get("fno") or [""])[0]
        did2=(q.get("did") or [did])[0]
        return detail,{
            "year":y,"month":m,"role":role,"text":txt,
            "url":dl,"download_url":dl,"parent":detail,
            "attachment_id":f"{did2}|{fid}",
        }

    if details:
        with ThreadPoolExecutor(max_workers=12) as pool:
            fm={pool.submit(inspect,item):item for item in details}
            for fut in as_completed(fm):
                item=fm[fut]
                try:
                    detail,att=fut.result();pages.append(detail)
                    if att:files.append(att)
                except Exception as e:
                    errors.append(f"detail {item[4]}: {type(e).__name__}: {e}")

    uniq={x["attachment_id"]:x for x in files}
    return sorted(uniq.values(),key=lambda x:(x["year"],x["month"],x["role"])),list(dict.fromkeys(pages)),errors

def parse_xlsx(att):
    try:
        blob=fetch(att["url"],True,att.get("parent",""))
        wb=openpyxl.load_workbook(io.BytesIO(blob),data_only=True,read_only=True)
    except Exception as e:
        return [],f"xlsx {att['url']}: {type(e).__name__}: {e}"
    rows=[]
    for ws in wb.worksheets:
        header_row=None;headers=[]
        for ri,row in enumerate(ws.iter_rows(min_row=1,max_row=min(ws.max_row,12),values_only=True),1):
            vals=[clean(v).replace(" ","") for v in row]
            joined="|".join(vals)
            if "사용일자" in joined and "사용처" in joined and "집행금액" in joined:
                header_row=ri;headers=vals;break
        if not header_row:continue
        def col(*names):
            for i,h in enumerate(headers):
                if any(n.replace(" ","") in h for n in names): return i
            return None
        di=col("사용일자"); pi=col("집행내역","목적"); mi=col("사용처"); ti=col("집행대상자")
        payi=col("집행구분"); peoi=col("인원"); ai=col("집행금액")
        if None in (di,pi,mi,ai):continue
        for ri,row in enumerate(ws.iter_rows(min_row=header_row+1,values_only=True),header_row+1):
            vals=list(row)
            def get(i): return vals[i] if i is not None and i<len(vals) else ""
            d=norm_date(get(di)); merchant=merchant_clean(get(mi)); amount=to_int(get(ai))
            if not d or not merchant or merchant in {"합계","계","사용처(장소)"} or amount is None:continue
            rows.append({
                "source_key":KEY,"institution":INSTITUTION,"cohort":"public_enterprise_leadership",
                "role":att["role"],"department":"","used_date":d,"used_time":"",
                "merchant":merchant,"address":"","purpose":clean(get(pi)),
                "target":clean(get(ti)),"payment_method":clean(get(payi)),
                "people":to_int(get(peoi)),"amount":amount,"source_amount_scale":1,
                "source_category":att["role"],
                "source_url":att["url"],"source_sheet":ws.title,"source_row":ri,
                "row_id":f"kca:{att['attachment_id']}:{ws.title}:{ri}",
            })
    return rows,None

def discover(year:int):
    attachments,pages,errors=discover_files(year);rows=[]
    if attachments:
        with ThreadPoolExecutor(max_workers=min(8,len(attachments))) as pool:
            fm={pool.submit(parse_xlsx,a):a for a in attachments}
            for fut in as_completed(fm):
                parsed,err=fut.result();rows.extend(parsed)
                if err:errors.append(err)
    rows.sort(key=lambda r:(r["used_date"],r["role"],r["row_id"]))
    return {
        "key":KEY,"institution":INSTITUTION,"cohort":"public_enterprise_leadership",
        "default_role":"기관장·부기관장·상임이사·상임위원·분쟁조정위원장·안전센터소장",
        "years":[year-1,year],"pages":pages,"attachments":attachments,
        "inline_rows":rows,"inline_replace":True,"errors":errors,
        "parseable_attachments":len(attachments),
        "status":"PARSEABLE_FOUND" if rows else ("FETCH_FAILED" if errors else "NO_FILES_FOUND"),
    }

def main():
    payload=json.loads(REPORT.read_text(encoding="utf-8"))
    fresh=discover(int(payload.get("year") or datetime.now().year))
    payload["sources"]=[x for x in payload.get("sources",[]) if x.get("key")!=KEY]+[fresh]
    REPORT.write_text(json.dumps(payload,ensure_ascii=False,indent=2),encoding="utf-8")
    print(json.dumps({
        "key":KEY,"status":fresh["status"],"attachments":len(fresh["attachments"]),
        "rows":len(fresh["inline_rows"]),"merchants":len({r["merchant"] for r in fresh["inline_rows"]}),
        "roles":sorted({r["role"] for r in fresh["inline_rows"]}),
        "date_min":min((r["used_date"] for r in fresh["inline_rows"]),default=""),
        "date_max":max((r["used_date"] for r in fresh["inline_rows"]),default=""),
        "errors":fresh["errors"][:8],
    },ensure_ascii=False))

if __name__=="__main__":main()
