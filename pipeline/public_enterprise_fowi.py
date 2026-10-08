#!/usr/bin/env python3
from __future__ import annotations

import html, io, json, re, time, urllib.parse, urllib.request
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from pathlib import Path
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
REPORT=ROOT/"reports"/"public-enterprise-discovery.json"
KEY="fowi"
INSTITUTION="한국산림복지진흥원"
LISTING="https://www.fowi.or.kr/user/publication/coreView.do?coreId=22&menu=core_"
BASE="https://www.fowi.or.kr"
UA="Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 Chrome/153 Safari/537.36"

def fetch(url:str,binary:bool=False,referer:str=""):
    last=None
    for attempt in range(3):
        try:
            req=urllib.request.Request(url,headers={
                "User-Agent":UA,"Accept":"*/*","Accept-Language":"ko-KR,ko;q=0.9",
                "Referer":referer or LISTING,"Connection":"close",
                "X-Requested-With":"XMLHttpRequest",
            })
            with urllib.request.urlopen(req,timeout=20) as r:
                raw=r.read()
                return raw if binary else raw.decode(r.headers.get_content_charset() or "utf-8","replace")
        except Exception as e:
            last=e
            if attempt<2:time.sleep(1+attempt)
    raise last

def clean(v):
    return " ".join(str(v or "").replace("\n"," ").split()).strip()

def date_norm(v):
    s=clean(v)
    m=re.search(r"(20\d{2})[-./](\d{1,2})[-./](\d{1,2})",s)
    if not m:return ""
    try:return datetime(*map(int,m.groups())).date().isoformat()
    except ValueError:return ""

def time_norm(v):
    s=clean(v)
    m=re.search(r"(\d{1,2}):(\d{2})(?::\d{2})?",s)
    if not m:return ""
    return f"{int(m.group(1)):02d}:{int(m.group(2)):02d}"

def to_int(v):
    s=re.sub(r"[^0-9.-]","",str(v or ""))
    try:return int(round(float(s)))
    except Exception:return None

def merchant_clean(v):
    s=clean(v)
    s=re.sub(r"\s*\(?☎?\)?\s*(?:0\d{1,3}|0507)[-\d]+\)?\s*$","",s)
    return s.strip()

def discover_files(year:int):
    years={year,year-1};errors=[];pages=[LISTING]
    try:doc=fetch(LISTING)
    except Exception as e:return [],pages,[f"listing {LISTING}: {type(e).__name__}: {e}"]
    candidates=[]
    for m in re.finditer(
        r"<a\b[^>]*onclick=['\"]fn_generateLink\(['\"](\d+)['\"]\);?['\"][^>]*>(.*?)</a>",
        doc,re.I|re.S
    ):
        fid,body=m.groups()
        title=clean(html.unescape(re.sub(r"<[^>]+>"," ",body)))
        ym=re.search(r"(20\d{2})년도(?:\s*(\d{1,2})월)?\s*기관장\s*및\s*임원\s*업무추진비\s*집행내역",title)
        if not ym:continue
        y=int(ym.group(1))
        if y not in years:continue
        reg=re.search(r"등록일\s*(20\d{2})[.\-](\d{1,2})[.\-](\d{1,2})",title)
        regkey=tuple(map(int,reg.groups())) if reg else (y,int(ym.group(2) or 0),0)
        candidates.append({"year":y,"file_id":fid,"text":title,"regkey":regkey})
    selected=[]
    for y in sorted(years):
        yy=[x for x in candidates if x["year"]==y]
        if not yy:continue
        x=max(yy,key=lambda a:(a["regkey"],int(a["file_id"])))
        try:
            link_json=json.loads(fetch(BASE+"/selectGenerateDownloadLink.do?"+urllib.parse.urlencode({"fileId":x["file_id"]})))
            rel=link_json.get("result","")
            if not rel:raise RuntimeError("empty generated download link")
            dl=urllib.parse.urljoin(BASE,rel)
        except Exception as e:
            errors.append(f"resolve {x['file_id']}: {type(e).__name__}: {e}");continue
        selected.append({
            "year":y,"file_id":x["file_id"],"text":x["text"],
            "url":dl,"download_url":dl,"parent":LISTING,
            "attachment_id":x["file_id"],
        })
    return selected,pages,errors

def parse_xlsx(att):
    try:
        blob=fetch(att["url"],True,att.get("parent",""))
        wb=load_workbook(io.BytesIO(blob),read_only=True,data_only=True)
    except Exception as e:return [],f"xlsx {att['file_id']}: {type(e).__name__}: {e}"
    rows=[]
    for ws in wb.worksheets:
        vals_all=[list(row) for row in ws.iter_rows(values_only=True)]
        if not vals_all:continue
        header=None
        for ri,vals in enumerate(vals_all):
            texts=[clean(v) for v in vals]
            if "사용일자" in texts and "사용처(장소)" in texts:
                header={}
                aliases=("사용일자","사용시간","직위","성명","사용목적(내역)","사용처(장소)","사용방법","대상인원(명)","집행금액(원)")
                for name in aliases:
                    if name in texts:header[name]=texts.index(name)
                continue
            if not header:continue
            def get(name):
                i=header.get(name)
                return vals[i] if i is not None and i<len(vals) else ""
            d=date_norm(get("사용일자"))
            merchant=merchant_clean(get("사용처(장소)"))
            if not d or not merchant or merchant in {"-","해당사항 없음","사용처(장소)"}:continue
            role=clean(get("직위")) or re.sub(r"^\d+월\s*","",ws.title)
            name=clean(get("성명"))
            rows.append({
                "source_key":KEY,"institution":INSTITUTION,
                "cohort":"public_enterprise_leadership",
                "role":role,"department":"","used_date":d,
                "used_time":time_norm(get("사용시간")),
                "merchant":merchant,"address":"",
                "purpose":clean(get("사용목적(내역)")),
                "target":"","payment_method":clean(get("사용방법")),
                "people":to_int(get("대상인원(명)")),"amount":to_int(get("집행금액(원)")),
                "source_amount_scale":1,
                "source_category":name or role,
                "source_url":att["url"],"source_sheet":ws.title,"source_row":ri+1,
                "row_id":f"fowi:{att['file_id']}:{ws.title}:{ri+1}",
            })
    return rows,None

def discover(year:int):
    attachments,pages,errors=discover_files(year);rows=[]
    with ThreadPoolExecutor(max_workers=max(1,min(2,len(attachments)))) as pool:
        fm={pool.submit(parse_xlsx,a):a for a in attachments}
        for fut in as_completed(fm):
            parsed,err=fut.result();rows.extend(parsed)
            if err:errors.append(err)
    rows.sort(key=lambda r:(r["used_date"],r["role"],r["row_id"]))
    return {
        "key":KEY,"institution":INSTITUTION,"cohort":"public_enterprise_leadership",
        "default_role":"진흥원장·부원장·국립산림치유원장",
        "years":[year-1,year],"pages":pages,"attachments":attachments,
        "inline_rows":rows,"inline_replace":True,"errors":errors,
        "parseable_attachments":len(attachments),
        "status":"PARSEABLE_FOUND" if rows else ("FETCH_FAILED" if errors else "NO_FILES_FOUND"),
    }

def main():
    payload=json.loads(REPORT.read_text(encoding="utf-8"))
    fresh=discover(int(payload.get("year") or datetime.now().year))
    payload["sources"]=[x for x in payload.get("sources",[]) if x.get("key")!=KEY]+[fresh]
    REPORT.write_text(json.dumps(payload,ensure_ascii=False,indent=2),encoding="utf-8")
    print(json.dumps({
        "key":KEY,"status":fresh["status"],"attachments":len(fresh["attachments"]),
        "rows":len(fresh["inline_rows"]),"merchants":len({r["merchant"] for r in fresh["inline_rows"]}),
        "roles":sorted({r["role"] for r in fresh["inline_rows"]}),
        "date_min":min((r["used_date"] for r in fresh["inline_rows"]),default=""),
        "date_max":max((r["used_date"] for r in fresh["inline_rows"]),default=""),
        "errors":fresh["errors"][:8],
    },ensure_ascii=False))

if __name__=="__main__":main()
